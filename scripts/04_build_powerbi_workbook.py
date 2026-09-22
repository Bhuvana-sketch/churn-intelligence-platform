"""
Stage 4: Assemble the Power BI-ready data model as a multi-sheet Excel workbook.
This is the file you'd point Power BI Desktop's "Get Data > Excel" at, or use
directly as the star-schema source (Fact_Customers + dimension/summary sheets).
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

import os
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "data")
DOCS = os.path.join(ROOT, "docs")
os.makedirs(DOCS, exist_ok=True)
WB_PATH = os.path.join(DOCS, "Churn_Intelligence_DataModel.xlsx")

df = pd.read_csv(f"{OUT}/telco_scored.csv")

fact_cols = [
    'customerID', 'gender', 'SeniorCitizen', 'Partner', 'Dependents', 'tenure',
    'Contract', 'PaperlessBilling', 'PaymentMethod', 'InternetService',
    'MonthlyCharges', 'TotalCharges', 'ServicesAdopted', 'EngagementScore',
    'RiskScore', 'PredictedChurnProbability', 'CohortName', 'Churn', 'ChurnFlag'
]
fact = df[fact_cols].rename(columns={'ChurnFlag': 'Churned'})

cohort_summary = pd.read_csv(f"{OUT}/cohort_summary.csv")
kpis = pd.read_csv(f"{OUT}/pg1_kpis.csv")
by_contract = pd.read_csv(f"{OUT}/pg3_by_contract.csv")
by_internet = pd.read_csv(f"{OUT}/pg3_by_internet.csv")
by_payment = pd.read_csv(f"{OUT}/pg3_by_payment.csv")
revenue_risk = pd.read_csv(f"{OUT}/pg4_revenue_at_risk.csv")
watchlist = pd.read_csv(f"{OUT}/pg5_watchlist_top100.csv")
feat_imp = pd.read_csv(f"{OUT}/feature_importance.csv")

with pd.ExcelWriter(WB_PATH, engine='openpyxl') as writer:
    fact.to_excel(writer, sheet_name='Fact_Customers', index=False)
    cohort_summary.to_excel(writer, sheet_name='Dim_CohortSummary', index=False)
    kpis.to_excel(writer, sheet_name='Summary_KPIs', index=False)
    by_contract.to_excel(writer, sheet_name='Summary_ByContract', index=False)
    by_internet.to_excel(writer, sheet_name='Summary_ByInternet', index=False)
    by_payment.to_excel(writer, sheet_name='Summary_ByPayment', index=False)
    revenue_risk.to_excel(writer, sheet_name='Summary_RevenueAtRisk', index=False)
    feat_imp.to_excel(writer, sheet_name='Summary_FeatureImportance', index=False)
    watchlist.to_excel(writer, sheet_name='Watchlist_Top100', index=False)

# ---------------- Formatting pass ----------------
wb = load_workbook(WB_PATH)

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    max_col = ws.max_column
    max_row = ws.max_row

    # Header row styling
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # Body styling + column width autosize
    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = len(str(ws.cell(row=1, column=col).value or ""))
        for row in range(2, max_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = BODY_FONT
            cell.border = BORDER
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
            # Percent-like columns
            header = str(ws.cell(row=1, column=col).value or "")
            if any(k in header for k in ["Rate", "Probability", "Pct"]) and isinstance(val, (int, float)):
                pass  # values already stored as 0-100 scale in these sheets; keep as number
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    ws.freeze_panes = "A2"
    if max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

wb.save(WB_PATH)
print(f"Saved formatted workbook -> {WB_PATH}")
print("Sheets:", wb.sheetnames)
